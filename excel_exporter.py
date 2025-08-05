"""Excel export functionality for enriched person data."""

from pathlib import Path
from typing import List
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from colorama import Fore, Style
from data_enricher import PersonInfo
from config import Config

class ExcelExporter:
    """Handles exporting person data to Excel files."""
    
    def __init__(self):
        """Initialize the Excel exporter."""
        pass
    
    def export_to_excel(self, person_data: List[PersonInfo], output_file: Path = None) -> Path:
        """Export person data to an Excel file with formatting."""
        if output_file is None:
            output_file = Config.OUTPUT_DIRECTORY / Config.DEFAULT_OUTPUT_FILE
        
        # Convert person data to DataFrame
        data_dicts = [person.to_dict() for person in person_data]
        df = pd.DataFrame(data_dicts)
        
        # Create Excel file with styling
        wb = Workbook()
        ws = wb.active
        ws.title = "Extracted Names"
        
        # Add header row with styling
        headers = list(df.columns)
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Add data rows
        for row_idx, row_data in enumerate(dataframe_to_rows(df, index=False, header=False), 2):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                
                # Wrap text for better readability
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                
                # Alternate row colors
                if row_idx % 2 == 0:
                    cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        
        # Adjust column widths
        column_widths = {
            'Name': 25,
            'Wikipedia Summary': 50,
            'Wikipedia URL': 40,
            'Google Results': 60,
            'AI Analysis': 80,
            'Known Details': 40
        }
        
        for col_idx, header in enumerate(headers, 1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = column_widths.get(header, 30)
        
        # Freeze the header row
        ws.freeze_panes = "A2"
        
        # Add summary sheet
        self._add_summary_sheet(wb, person_data)
        
        # Save the file
        wb.save(output_file)
        
        print(f"{Fore.GREEN}Excel file exported to: {output_file}{Style.RESET_ALL}")
        print(f"Total names processed: {len(person_data)}")
        
        return output_file
    
    def _add_summary_sheet(self, workbook: Workbook, person_data: List[PersonInfo]):
        """Add a summary sheet with statistics."""
        ws = workbook.create_sheet(title="Summary")
        
        # Statistics
        total_names = len(person_data)
        names_with_wikipedia = sum(1 for p in person_data if p.wikipedia_summary)
        names_with_google = sum(1 for p in person_data if p.google_results)
        names_with_ai = sum(1 for p in person_data if p.ai_analysis and "unavailable" not in p.ai_analysis.lower())
        
        # Add summary data
        summary_data = [
            ["Statistic", "Value"],
            ["Total Names Extracted", total_names],
            ["Names with Wikipedia Info", names_with_wikipedia],
            ["Names with Google Results", names_with_google],
            ["Names with AI Analysis", names_with_ai],
            ["Success Rate - Wikipedia", f"{(names_with_wikipedia/total_names)*100:.1f}%" if total_names > 0 else "0%"],
            ["Success Rate - Google", f"{(names_with_google/total_names)*100:.1f}%" if total_names > 0 else "0%"],
            ["Success Rate - AI", f"{(names_with_ai/total_names)*100:.1f}%" if total_names > 0 else "0%"]
        ]
        
        # Add data to worksheet
        for row_idx, (label, value) in enumerate(summary_data, 1):
            ws.cell(row=row_idx, column=1, value=label)
            ws.cell(row=row_idx, column=2, value=value)
            
            # Style header row
            if row_idx == 1:
                ws.cell(row=row_idx, column=1).font = Font(bold=True)
                ws.cell(row=row_idx, column=2).font = Font(bold=True)
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 20
        
        # Add top names by information richness
        ws.cell(row=len(summary_data) + 3, column=1, value="Most Complete Profiles:").font = Font(bold=True)
        
        # Sort people by completeness score
        scored_people = []
        for person in person_data:
            score = 0
            if person.wikipedia_summary: score += 3
            if person.google_results: score += 2
            if person.ai_analysis and "unavailable" not in person.ai_analysis.lower(): score += 2
            if person.known_details: score += 1
            scored_people.append((person.name, score))
        
        scored_people.sort(key=lambda x: x[1], reverse=True)
        
        start_row = len(summary_data) + 4
        for idx, (name, score) in enumerate(scored_people[:10], start_row):  # Top 10
            ws.cell(row=idx, column=1, value=name)
            ws.cell(row=idx, column=2, value=f"Score: {score}/8")
    
    def export_simple_list(self, names: List[str], output_file: Path = None) -> Path:
        """Export a simple list of names to Excel (no enrichment)."""
        if output_file is None:
            output_file = Config.OUTPUT_DIRECTORY / "simple_names_list.xlsx"
        
        # Create simple DataFrame
        df = pd.DataFrame({'Names': sorted(names)})
        
        # Save to Excel
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Names', index=False)
            
            # Format the worksheet
            workbook = writer.book
            worksheet = writer.sheets['Names']
            
            # Style header
            header_cell = worksheet['A1']
            header_cell.font = Font(bold=True, color="FFFFFF")
            header_cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_cell.alignment = Alignment(horizontal="center")
            
            # Adjust column width
            worksheet.column_dimensions['A'].width = 30
        
        print(f"{Fore.GREEN}Simple names list exported to: {output_file}{Style.RESET_ALL}")
        print(f"Total names: {len(names)}")
        
        return output_file
    
    def export_csv(self, person_data: List[PersonInfo], output_file: Path = None) -> Path:
        """Export person data to CSV format."""
        if output_file is None:
            output_file = Config.OUTPUT_DIRECTORY / "extracted_names.csv"
        
        # Convert to DataFrame and save as CSV
        data_dicts = [person.to_dict() for person in person_data]
        df = pd.DataFrame(data_dicts)
        df.to_csv(output_file, index=False, encoding='utf-8')
        
        print(f"{Fore.GREEN}CSV file exported to: {output_file}{Style.RESET_ALL}")
        return output_file